from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from groq import Groq
from dotenv import load_dotenv
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGB
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
from reportlab.lib import colors
from pypdf import PdfReader
from duckduckgo_search import DDGS
import os
import io
import re
import requests

load_dotenv()

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

client = Groq(api_key=os.getenv("GROQ_API_KEY"))

BASE_PROMPT = "Tu es Nexa AI, l'assistant intelligent créé par Nexora, fondée par Randriafenosoa Mamiratiniaina Kevin. Tu réponds TOUJOURS dans la même langue que l'utilisateur. Tu ne mentionnes jamais OpenAI ni ChatGPT. Utilise le format Markdown pour structurer tes réponses : titres avec ##, listes avec -, tableaux avec |."


def demander(prompt, system=None, model="openai/gpt-oss-120b"):
    messages = [{"role": "system", "content": system or BASE_PROMPT}]
    messages.append({"role": "user", "content": prompt})
    response = client.chat.completions.create(
        messages=messages,
        model=model,
    )
    return response.choices[0].message.content


def rechercher_web(query):
    try:
        with DDGS() as ddgs:
            resultats = list(ddgs.text(query, max_results=5))
        if not resultats:
            return ""
        texte = "Résultats de recherche récents :\n\n"
        for i, r in enumerate(resultats, 1):
            texte += f"{i}. {r.get('title', '')}\n{r.get('body', '')}\nSource : {r.get('href', '')}\n\n"
        return texte
    except Exception:
        return ""


@app.get("/")
def racine():
    return {"message": "Nexa AI API fonctionne !"}


# ============ CHAT ============
@app.post("/chat")
def chat(data: dict):
    message = data.get("message", "")
    historique = data.get("historique", [])
    persona = data.get("persona", "")
    recherche = data.get("recherche", False)

    system = BASE_PROMPT
    if persona:
        system += f" Comportement souhaité : {persona}."

    if recherche:
        contexte = rechercher_web(message)
        if contexte:
            system += "\n\nUtilise ces résultats de recherche pour répondre :\n" + contexte

    messages = [{"role": "system", "content": system}]
    for msg in historique[-20:]:
        role = "user" if msg["type"] == "user" else "assistant"
        messages.append({"role": role, "content": msg["texte"]})
    messages.append({"role": "user", "content": message})

    response = client.chat.completions.create(
        messages=messages,
        model="openai/gpt-oss-120b",
    )
    return {"reply": response.choices[0].message.content}


# ============ QUIZ ============
@app.post("/quiz")
def quiz(data: dict):
    sujet = data.get("sujet", "")
    nb = data.get("nb", 5)
    prompt = f"""Crée un quiz de {nb} questions à choix multiples (QCM) sur : {sujet}.

FORMAT STRICT :
Question 1 : [question]
A) [option]
B) [option]
C) [option]
D) [option]
Réponse correcte : [lettre]

(répète pour chaque question, sans autre texte)"""
    contenu = demander(prompt, "Tu crées des quiz QCM en respectant STRICTEMENT le format.")
    return {"reply": contenu, "quiz": contenu}


# ============ FICHE ============
@app.post("/fiche")
def fiche(data: dict):
    sujet = data.get("sujet", "")
    prompt = f"""Crée une fiche de révision complète et structurée sur : {sujet}.

Utilise le format Markdown :
# Titre
## Section
- Point 1
- Point 2

**Définitions importantes** en gras.
Tableaux si nécessaire."""
    return {"reply": demander(prompt, "Tu crées des fiches de révision claires et bien structurées en Markdown.")}


# ============ CV ============
@app.post("/cv")
def cv(data: dict):
    infos = data.get("infos", "")
    prompt = f"Crée un CV professionnel moderne basé sur : {infos}. Utilise le Markdown."
    return {"reply": demander(prompt, "Tu es un expert en recrutement.")}


# ============ CORRECTION ============
@app.post("/corriger")
def corriger(data: dict):
    texte = data.get("texte", "")
    prompt = f"Corrige ce devoir :\n\n{texte}\n\nDonne : fautes, améliorations, note sur 20, commentaire."
    return {"reply": demander(prompt, "Tu corriges avec bienveillance.")}


# ============ UTILITAIRES DE PARSING ============
def extraire_tableaux_markdown(texte):
    """Extrait les tableaux Markdown du texte"""
    lignes = texte.split('\n')
    tableaux = []
    tableau_actuel = []
    
    for ligne in lignes:
        if '|' in ligne and ligne.strip().startswith('|'):
            tableau_actuel.append(ligne)
        else:
            if tableau_actuel:
                tableaux.append(tableau_actuel)
                tableau_actuel = []
    if tableau_actuel:
        tableaux.append(tableau_actuel)
    
    return tableaux


def parser_tableau(tableau_lines):
    """Parse un tableau Markdown en liste de listes"""
    rows = []
    for line in tableau_lines:
        if re.match(r'^\|[\s\-:|]+\|$', line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append(cells)
    return rows


# ============ GÉNÉRATION WORD ============
@app.post("/word")
def generer_word(data: dict):
    sujet = data.get("sujet", "")
    
    prompt = f"""Rédige un document complet et professionnel sur : {sujet}

Structure le document avec :
- Un titre principal (# Titre)
- Des sections (## Section)
- Des sous-sections (### Sous-section)
- Des paragraphes explicatifs
- Des listes à puces (- point)
- Des tableaux si pertinent (| Colonne | Colonne |)

Format Markdown."""

    contenu = demander(prompt, "Tu es un rédacteur professionnel. Tu structures tes documents en Markdown avec titres, listes et tableaux.")

    doc = Document()
    
    # Styles
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Titre principal
    doc.add_heading(sujet, 0)
    
    # Traiter ligne par ligne
    lignes = contenu.split('\n')
    i = 0
    while i < len(lignes):
        ligne = lignes[i].strip()
        
        if not ligne:
            i += 1
            continue
        
        # Tableau Markdown
        if ligne.startswith('|') and '|' in ligne:
            tableau_lines = []
            while i < len(lignes) and lignes[i].strip().startswith('|'):
                tableau_lines.append(lignes[i].strip())
                i += 1
            
            rows = parser_tableau(tableau_lines)
            if rows:
                nb_cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=nb_cols)
                table.style = 'Light Grid Accent 1'
                
                for ri, row in enumerate(rows):
                    for ci in range(nb_cols):
                        cell = table.cell(ri, ci)
                        cell.text = row[ci] if ci < len(row) else ''
                        if ri == 0:
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.bold = True
            continue
        
        # Titres
        if ligne.startswith('# '):
            doc.add_heading(ligne[2:].strip(), level=1)
        elif ligne.startswith('## '):
            doc.add_heading(ligne[3:].strip(), level=2)
        elif ligne.startswith('### '):
            doc.add_heading(ligne[4:].strip(), level=3)
        # Listes à puces
        elif ligne.startswith('- ') or ligne.startswith('* '):
            p = doc.add_paragraph(ligne[2:].strip(), style='List Bullet')
        # Listes numérotées
        elif re.match(r'^\d+\.\s', ligne):
            texte_item = re.sub(r'^\d+\.\s', '', ligne)
            doc.add_paragraph(texte_item, style='List Number')
        # Citation
        elif ligne.startswith('> '):
            p = doc.add_paragraph(ligne[2:].strip())
            p.italic = True
        # Paragraphe normal
        else:
            # Gérer le gras **texte**
            p = doc.add_paragraph()
            parties = re.split(r'(\*\*[^*]+\*\*)', ligne)
            for partie in parties:
                if partie.startswith('**') and partie.endswith('**'):
                    run = p.add_run(partie[2:-2])
                    run.bold = True
                else:
                    p.add_run(partie)
        
        i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=nexa_document.docx"}
    )


# ============ GÉNÉRATION EXCEL ============
@app.post("/excel")
def generer_excel(data: dict):
    sujet = data.get("sujet", "")
    
    prompt = f"""Crée un tableau de données structuré sur : {sujet}

FORMAT STRICT : tableau Markdown uniquement, avec en-têtes.

| Colonne 1 | Colonne 2 | Colonne 3 |
|-----------|-----------|-----------|
| valeur    | valeur    | valeur    |

Donne au moins 5 lignes de données réalistes. Aucun autre texte."""

    contenu = demander(prompt, "Tu crées des tableaux de données. Tu réponds UNIQUEMENT avec un tableau Markdown.")

    wb = Workbook()
    ws = wb.active
    ws.title = (sujet[:30] or "Données")

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    lignes = contenu.split('\n')
    row_index = 1
    
    for ligne in lignes:
        ligne = ligne.strip()
        if not ligne or not ligne.startswith('|'):
            continue
        if re.match(r'^\|[\s\-:|]+\|$', ligne):
            continue
        
        cellules = [c.strip() for c in ligne.strip('|').split('|')]
        for ci, val in enumerate(cellules, 1):
            cell = ws.cell(row=row_index, column=ci, value=val)
            cell.border = border
            if row_index == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        
        row_index += 1

    # Ajuster la largeur des colonnes
    for col in ws.columns:
        max_length = 0
        column_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 4, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=nexa_tableau.xlsx"}
    )


# ============ GÉNÉRATION POWERPOINT ============
@app.post("/pptx")
def generer_pptx(data: dict):
    sujet = data.get("sujet", "")
    
    prompt = f"""Crée une présentation PowerPoint professionnelle sur : {sujet}

FORMAT STRICT :
Slide 1 : [Titre de la slide]
- [point 1]
- [point 2]
- [point 3]

Slide 2 : [Titre de la slide]
- [point 1]
- [point 2]

Continue pour 6 slides au total. Aucun autre texte."""

    contenu = demander(prompt, "Tu crées des présentations structurées. Tu respectes STRICTEMENT le format demandé.")

    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)

    slides_data = []
    slide_actuelle = None

    for ligne in contenu.split("\n"):
        ligne = ligne.strip()
        if not ligne:
            continue
        if ligne.lower().startswith("slide"):
            if slide_actuelle:
                slides_data.append(slide_actuelle)
            titre = ligne.split(":", 1)[1].strip() if ":" in ligne else ligne
            slide_actuelle = {"titre": titre, "contenu": []}
        elif ligne.startswith("- ") and slide_actuelle:
            slide_actuelle["contenu"].append(ligne[2:])

    if slide_actuelle:
        slides_data.append(slide_actuelle)

    if not slides_data:
        slides_data = [{"titre": sujet, "contenu": [contenu]}]

    # Slide de titre
    slide_titre = prs.slides.add_slide(prs.slide_layouts[0])
    slide_titre.shapes.title.text = sujet
    if len(slide_titre.placeholders) > 1:
        slide_titre.placeholders[1].text = "Présentation générée par Nexa AI"

    # Slides de contenu
    for s in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s["titre"]
        if len(slide.placeholders) > 1:
            tf = slide.placeholders[1].text_frame
            tf.text = s["contenu"][0] if s["contenu"] else ""
            for point in s["contenu"][1:]:
                p = tf.add_paragraph()
                p.text = point
                p.level = 0

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": "attachment; filename=nexa_presentation.pptx"}
    )


# ============ GÉNÉRATION PDF ============
@app.post("/pdf")
def generer_pdf(data: dict):
    sujet = data.get("sujet", "")
    
    prompt = f"""Rédige un document professionnel sur : {sujet}

Structure avec :
# Titre
## Section
Paragraphes et listes.

Format Markdown."""

    contenu = demander(prompt, "Tu rédiges des documents professionnels en Markdown.")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    
    # Styles personnalisés
    titre_style = ParagraphStyle(
        'Titre',
        parent=styles['Heading1'],
        fontSize=22,
        textColor=colors.HexColor('#667EEA'),
        spaceAfter=20,
        alignment=1
    )
    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Heading2'],
        fontSize=16,
        textColor=colors.HexColor('#333333'),
        spaceAfter=12,
        spaceBefore=15
    )
    h3_style = ParagraphStyle(
        'H3',
        parent=styles['Heading3'],
        fontSize=13,
        textColor=colors.HexColor('#555555'),
        spaceAfter=8,
        spaceBefore=10
    )
    normal_style = ParagraphStyle(
        'Normal',
        parent=styles['Normal'],
        fontSize=11,
        leading=16,
        spaceAfter=8
    )

    story = [Paragraph(sujet, titre_style), Spacer(1, 0.5*cm)]

    for ligne in contenu.split('\n'):
        ligne = ligne.strip()
        if not ligne:
            continue
        
        if ligne.startswith('# '):
            story.append(Paragraph(ligne[2:], titre_style))
        elif ligne.startswith('## '):
            story.append(Paragraph(ligne[3:], h2_style))
        elif ligne.startswith('### '):
            story.append(Paragraph(ligne[4:], h3_style))
        elif ligne.startswith('- ') or ligne.startswith('* '):
            story.append(Paragraph('• ' + ligne[2:], normal_style))
        elif re.match(r'^\d+\.\s', ligne):
            story.append(Paragraph(ligne, normal_style))
        else:
            # Gras **texte**
            texte = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', ligne)
            story.append(Paragraph(texte, normal_style))

    doc.build(story)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=nexa_document.pdf"}
    )


# ============ GÉNÉRATION IMAGE ============
@app.post("/image")
def generer_image(data: dict):
    prompt = data.get("prompt", "")
    prompt_encode = requests.utils.quote(prompt)
    url = f"https://image.pollinations.ai/prompt/{prompt_encode}?width=1024&height=1024&nologo=true"
    return {"url": url}


# ============ ANALYSE FICHIERS ============
@app.post("/analyser")
async def analyser(file: UploadFile = File(...)):
    contenu = await file.read()
    nom = file.filename.lower()

    try:
        if nom.endswith(".docx"):
            doc = Document(io.BytesIO(contenu))
            texte = "\n".join([p.text for p in doc.paragraphs])
        elif nom.endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(contenu))
            ws = wb.active
            lignes = []
            for row in ws.iter_rows(values_only=True):
                lignes.append(", ".join([str(c) if c is not None else "" for c in row]))
            texte = "\n".join(lignes)
        elif nom.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(contenu))
            texte = "\n".join([p.extract_text() for p in reader.pages])
        else:
            texte = contenu.decode("utf-8", errors="ignore")

        analyse = demander(f"Analyse ce document :\n\n{texte[:3000]}\n\nRésume-le, donne les points clés et des recommandations en Markdown.")
        return {"reply": analyse}
    except Exception as e:
        return {"reply": f"Erreur d'analyse : {str(e)}"}
